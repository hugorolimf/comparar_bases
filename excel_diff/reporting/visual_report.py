from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from excel_diff.models import ComparisonResult, DiffKeyPair, RowDiff


def write_visual_report(result: ComparisonResult, path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    write_sheet(workbook, "Igual", [row for row in result.matched_rows if row.status == "matched"], result, source_kind="base")
    write_sheet(workbook, "Alteração", [row for row in result.matched_rows if row.status == "changed"], result, source_kind="base")
    write_sheet(workbook, "Exclusão", result.only_in_base, result, source_kind="base")
    write_sheet(workbook, "Adição", result.only_in_compare, result, source_kind="compare")

    _write_guide_sheet(workbook, result)

    workbook.save(path)


def _total_columns(result: ComparisonResult, title: str) -> int:
    cols = 1 + len(result.base_profile.headers) + 2  # Categoria + headers + linhas
    if title == "Alteração":
        cols += 2 * len(result.diff_key_pairs)
    return cols


def _get_sheet_descriptions(title: str) -> list[str]:
    if title == "Igual":
        return [
            "Descrição: Esta planilha contém as linhas que são idênticas entre o arquivo base e o arquivo de comparação. "
            "Nenhuma diferença foi encontrada nos valores das colunas mapeadas.",
            "Significado: Os registros existem em ambos os arquivos e possuem os mesmos valores. As colunas 'Linha Base' e 'Linha Comparacao' indicam a posição original de cada registro nos respectivos arquivos.",
        ]
    if title == "Alteração":
        return [
            "Descrição: Esta planilha exibe os registros que existem em ambos os arquivos, mas apresentam diferenças nos valores das colunas de identificação selecionadas.",
            "Significado: As colunas 'ibase-{coluna}' mostram o valor original no arquivo base, enquanto 'diff-{coluna}' mostram o valor correspondente no arquivo de comparação. "
            "Isso permite identificar rapidamente quais campos foram modificados. As colunas 'Linha Base' e 'Linha Comparacao' indicam a posição original.",
        ]
    if title == "Exclusão":
        return [
            "Descrição: Esta planilha lista os registros que estão presentes apenas no arquivo base e não foram encontrados no arquivo de comparação.",
            "Significado: Essas linhas podem indicar registros removidos ou que não possuem correspondência pela chave escolhida. A coluna 'Linha Base' indica a posição original no arquivo base.",
        ]
    if title == "Adição":
        return [
            "Descrição: Esta planilha lista os registros que estão presentes apenas no arquivo de comparação e não existem no arquivo base.",
            "Significado: Essas linhas representam registros novos ou adicionados. Os valores são projetados para os nomes de colunas do arquivo base quando há mapeamento correspondente. "
            "A coluna 'Linha Comparacao' indica a posição original no arquivo de comparação.",
        ]
    return []


def write_sheet(workbook: Workbook, title: str, rows: list[RowDiff], result: ComparisonResult, source_kind: str) -> None:
    sheet = workbook.create_sheet(title)
    total_cols = _total_columns(result, title)

    # Linhas explicativas no topo
    for desc in _get_sheet_descriptions(title):
        sheet.append([desc])
        row_num = sheet.max_row
        sheet.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=total_cols)
        for cell in sheet[row_num]:
            cell.font = Font(italic=True, size=10, color="333333")
            cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[row_num].height = 30

    # Cabeçalhos de dados
    headers = ["Categoria"]
    headers.extend(result.base_profile.headers)
    if title == "Alteração":
        for pair in result.diff_key_pairs:
            headers.extend([f"ibase-{pair.base_column}", f"diff-{pair.base_column}"])
    headers.extend(["Linha Base", "Linha Comparacao"])
    sheet.append(headers)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    compare_to_base = {mapping.compare_column: mapping.base_column for mapping in result.column_mappings}

    for row in rows:
        base_values = resolve_base_values(row, result, source_kind, compare_to_base)
        values = [title]
        values.extend([base_values.get(column, "") for column in result.base_profile.headers])

        if title == "Alteração":
            for pair in result.diff_key_pairs:
                identifier = find_identifier(row.diff_identifiers, pair)
                values.extend([identifier.get("base_value", ""), identifier.get("compare_value", "")])

        values.extend([
            row.base_row_number if row.base_row_number is not None else "",
            row.compare_row_number if row.compare_row_number is not None else "",
        ])

        sheet.append(values)

    # Ajuste básico de largura das colunas
    from openpyxl.utils import get_column_letter

    for col_idx in range(1, total_cols + 1):
        column_letter = get_column_letter(col_idx)
        max_length = 0
        for cell in sheet[column_letter]:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 60)
        sheet.column_dimensions[column_letter].width = adjusted_width


def _write_guide_sheet(workbook: Workbook, result: ComparisonResult) -> None:
    sheet = workbook.create_sheet("Sobre", 0)

    # Título principal
    sheet.append(["Relatório Visual de Comparação"])
    sheet.merge_cells("A1:E1")
    for cell in sheet[1]:
        cell.font = Font(size=16, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.append([])

    # Informações gerais
    sheet.append(["Informações Gerais"])
    sheet.merge_cells("A3:E3")
    for cell in sheet[3]:
        cell.font = Font(size=12, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
    sheet.row_dimensions[3].height = 22

    sheet.append(["Arquivo Base", str(result.base_profile.path)])
    sheet.append(["Aba Base", result.base_profile.sheet_name])
    sheet.append(["Arquivo Comparação", str(result.compare_profile.path)])
    sheet.append(["Aba Comparação", result.compare_profile.sheet_name])
    sheet.append(["Coluna Chave Base", result.key_column])
    sheet.append(["Coluna Chave Comparação", result.resolved_compare_key])

    for row in sheet["4:9"]:
        for cell in row:
            cell.font = Font(size=10)
            if cell.column == 1:
                cell.font = Font(size=10, bold=True)

    sheet.append([])

    # Resumo das diferenças
    sheet.append(["Resumo das Diferenças"])
    sheet.merge_cells("A11:E11")
    for cell in sheet[11]:
        cell.font = Font(size=12, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
    sheet.row_dimensions[11].height = 22

    sheet.append(["Categoria", "Quantidade", "Descrição"])
    for cell in sheet[12]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    counts = {
        "Igual": len([r for r in result.matched_rows if r.status == "matched"]),
        "Alteração": len([r for r in result.matched_rows if r.status == "changed"]),
        "Exclusão": len(result.only_in_base),
        "Adição": len(result.only_in_compare),
    }

    descriptions = {
        "Igual": "Registros idênticos em ambos os arquivos.",
        "Alteração": "Registros com valores diferentes nas colunas de identificação.",
        "Exclusão": "Registros presentes apenas no arquivo base.",
        "Adição": "Registros presentes apenas no arquivo de comparação.",
    }

    row_idx = 13
    for cat in ["Igual", "Alteração", "Exclusão", "Adição"]:
        sheet.append([cat, counts[cat], descriptions[cat]])
        for cell in sheet[row_idx]:
            cell.font = Font(size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cat == "Alteração":
                cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            elif cat == "Exclusão":
                cell.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
            elif cat == "Adição":
                cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        row_idx += 1

    sheet.append([])

    # Descrição das planilhas
    sheet.append(["Descrição das Planilhas"])
    sheet.merge_cells(start_row=row_idx + 1, start_column=1, end_row=row_idx + 1, end_column=5)
    for cell in sheet[row_idx + 1]:
        cell.font = Font(size=12, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
    sheet.row_dimensions[row_idx + 1].height = 22

    row_idx += 2
    sheet_descriptions = [
        (
            "Igual",
            "Contém as linhas idênticas entre os dois arquivos. Nenhuma diferença foi encontrada nos valores das colunas mapeadas. "
            "As colunas 'Linha Base' e 'Linha Comparacao' indicam a posição original de cada registro.",
        ),
        (
            "Alteração",
            "Contém linhas que existem nos dois arquivos, mas com valores diferentes nas colunas de identificação. "
            "As colunas 'ibase-{coluna}' mostram o valor original no arquivo base, enquanto 'diff-{coluna}' mostram o valor no arquivo de comparação. "
            "Isso facilita a identificação rápida dos campos modificados.",
        ),
        (
            "Exclusão",
            "Contém linhas que só existem no arquivo base (não encontradas no arquivo de comparação pela chave escolhida). "
            "Essas linhas podem indicar registros removidos ou sem correspondência.",
        ),
        (
            "Adição",
            "Contém linhas que só existem no arquivo de comparação (não encontradas no arquivo base pela chave escolhida). "
            "Os valores são projetados para os nomes de colunas do arquivo base quando há mapeamento correspondente. "
            "Representam registros novos ou adicionados.",
        ),
    ]

    for name, desc in sheet_descriptions:
        sheet.append([name, desc])
        sheet.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=5)
        for cell in sheet[row_idx]:
            cell.font = Font(size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.column == 1:
                cell.font = Font(size=10, bold=True)
        row_idx += 1

    # Ajuste de larguras
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 70
    sheet.column_dimensions["D"].width = 15
    sheet.column_dimensions["E"].width = 15


def resolve_base_values(
    row: RowDiff,
    result: ComparisonResult,
    source_kind: str,
    compare_to_base: dict[str, str],
) -> dict[str, object]:
    if source_kind == "compare":
        return project_compare_to_base(row.compare_values, compare_to_base, result.base_profile.headers)
    return row.base_values


def project_compare_to_base(compare_values: dict[str, object], compare_to_base: dict[str, str], base_headers: list[str]) -> dict[str, object]:
    projected = {header: "" for header in base_headers}
    for compare_column, value in compare_values.items():
        base_column = compare_to_base.get(compare_column)
        if base_column in projected:
            projected[base_column] = value
    return projected


def find_identifier(diff_identifiers: list[dict[str, object]], pair: DiffKeyPair) -> dict[str, object]:
    for identifier in diff_identifiers:
        if identifier.get("base_column") == pair.base_column and identifier.get("compare_column") == pair.compare_column:
            return identifier
    return {}
